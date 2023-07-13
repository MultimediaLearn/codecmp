import openpyxl
from openpyxl import Workbook
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import (
    LineChart,
    Reference,
)
from openpyxl.chart.axis import DateAxis

wb = Workbook()
ws = wb.active
ws2 = wb.create_sheet("show")
img = openpyxl.drawing.image.Image("small.png")
ws2.add_image(img, "F10")

rows = [
    [400, 300, 250, 400],
    [400, 250, 300, 400],
    [500, 300, 450, 400],
    [300, 250, 400, 400],
]

for row in rows:
    ws.append(row)

c1 = LineChart()
c1.title = "Line Chart"
c1.style = 13

data = Reference(ws, min_col=2, min_row=1, max_col=4, max_row=4)
c1.add_data(data)
c1.legend = None


# Style the lines
s1 = c1.series[0]
s1.marker.symbol = "triangle"
s1.marker.graphicalProperties.solidFill = "FF0000" # Marker filling
s1.marker.graphicalProperties.line.solidFill = "FF0000" # Marker outline

s1.graphicalProperties.line.noFill = True

s2 = c1.series[1]
s2.graphicalProperties.line.solidFill = "00AAAA"
s2.graphicalProperties.line.dashStyle = "sysDot"
s2.graphicalProperties.line.width = 100050 # width in EMUs

s2 = c1.series[2]
s2.smooth = True # Make the line smooth
c1.width = 5
c1.height = 2

ws2.add_chart(c1, "A10")

wb.save("img.xlsx")
